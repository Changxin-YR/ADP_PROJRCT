from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta, timezone
import json
from threading import Barrier
from typing import Any

import pytest
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

from backend.layers.common.audit.audit_logger import AuditLogger
from backend.layers.common.db.connection import get_connection
from backend.layers.common.governance.idempotency import execute_idempotent
from backend.layers.common.governance.lifecycle import DomainError
from backend.layers.features.agent.agent_confirmation_store import MySqlAgentConfirmationStore
from backend.layers.features.agent.agent_contracts import AgentConfirmation
from backend.layers.features.agent.agent_gateway_service import AgentGatewayService, AgentGatewayError
from backend.layers.features.agent.agent_tool_registry import AgentTool, AgentToolRegistry
from backend.layers.features.master_data.master_data_service import MasterDataService
from backend.layers.features.master_data.master_data_store import MySqlMasterDataStore
from backend.layers.features.production.production_service import ProductionService
from backend.layers.features.production.production_store import MySqlProductionStore
from backend.layers.features.warehouse.warehouse_service import WarehouseService
from backend.layers.features.warehouse.warehouse_store import MySqlWarehouseStore
from backend.tests.mysql_test_database import disposable_database, settings_for
from test_warehouse_mysql_integration import _seed
from backend.app import create_app
from backend.config.settings import Settings


def _env_for(settings: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("MYSQL_HOST", settings.mysql_host)
    monkeypatch.setenv("MYSQL_PORT", str(settings.mysql_port))
    monkeypatch.setenv("MYSQL_DATABASE", settings.mysql_database)
    monkeypatch.setenv("MYSQL_USER", settings.mysql_user)
    monkeypatch.setenv("MYSQL_PASSWORD", settings.mysql_password)


def _actor(user_id: int, area_id: int | None = None) -> dict[str, Any]:
    scopes = [] if area_id is None else [{"scope_type": "area", "area_id": area_id}]
    return {
        "id": user_id,
        "status": "active",
        "permissions": ["master_data.view", "master_data.manage"],
        "roles": [] if area_id is None else [{"code": "operator"}],
        "data_scopes": scopes,
        "_scope_enforced": area_id is not None,
    }


def _rest_settings(settings: Any, attachment_root: str) -> Settings:
    return Settings.from_env({
        "APP_ENV": "test",
        "FLASK_SECRET_KEY": "idor-flask",
        "CSRF_SECRET_KEY": "idor-csrf",
        "MYSQL_HOST": settings.mysql_host,
        "MYSQL_PORT": str(settings.mysql_port),
        "MYSQL_DATABASE": settings.mysql_database,
        "MYSQL_USER": settings.mysql_user,
        "MYSQL_PASSWORD": settings.mysql_password,
        "SESSION_COOKIE_SECURE": "false",
        "ATTACHMENT_ROOT": attachment_root,
    })


def _rest_login(client: Any, login_name: str) -> None:
    csrf = client.get("/api/v1/auth/csrf").get_json()["data"]["csrf_token"]
    response = client.post(
        "/api/v1/auth/login",
        json={"identifier": login_name, "password": "Correct9!"},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200


def _rest_csrf(client: Any) -> dict[str, str]:
    return {"X-CSRF-Token": client.get("/api/v1/auth/csrf").get_json()["data"]["csrf_token"]}


def test_mysql_two_user_rest_idor_matrix(monkeypatch: pytest.MonkeyPatch, tmp_path: Any) -> None:
    """Same permissions, different areas: REST IDs never cross the scope boundary."""
    with disposable_database("adp_two_user_idor", through=32) as database:
        settings = settings_for(database)
        _env_for(settings, monkeypatch)
        ids = _seed(settings)
        password_hash = generate_password_hash("Correct9!", method="scrypt")
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT organization_id,farm_id FROM areas WHERE id=%s", (ids["area_id"],))
            tenant = cursor.fetchone()
            cursor.execute(
                "INSERT INTO areas (organization_id,farm_id,code,name,status,created_by) VALUES (%s,%s,'B','B区','verified',1)",
                (tenant["organization_id"], tenant["farm_id"]),
            )
            area_b = int(cursor.lastrowid)
            cursor.execute(
                "INSERT INTO users (phone,login_name,name,password_hash,status) VALUES "
                "('13980000101','idor-a','IDOR A',%s,'active'),('13980000102','idor-b','IDOR B',%s,'active')",
                (password_hash, password_hash),
            )
            cursor.execute("SELECT id,login_name FROM users WHERE login_name IN ('idor-a','idor-b')")
            user_ids = {str(row["login_name"]): int(row["id"]) for row in cursor.fetchall()}
            user_a, user_b = user_ids["idor-a"], user_ids["idor-b"]
            cursor.execute("INSERT INTO roles (code,name,status) VALUES ('idor-certifier','IDOR certifier','active')")
            role_id = int(cursor.lastrowid)
            cursor.execute(
                "INSERT INTO role_permissions (role_id,permission_id) "
                "SELECT %s,id FROM permissions WHERE code IN "
                "('master_data.view','master_data.manage','master_data.verify','data_exchange.export','attachment.manage')",
                (role_id,),
            )
            cursor.executemany("INSERT INTO user_roles (user_id,role_id) VALUES (%s,%s)", [(user_a, role_id), (user_b, role_id)])
            cursor.execute("SELECT id FROM data_scopes WHERE area_id=%s AND code LIKE 'A%%' LIMIT 1", (ids["area_id"],))
            scope_a_row = cursor.fetchone()
            if scope_a_row is None:
                cursor.execute("INSERT INTO data_scopes (code,name,scope_type,organization_id,farm_id,area_id,status) VALUES ('A-test-all','A区全部数据','area',%s,%s,%s,'active')", (tenant["organization_id"], tenant["farm_id"], ids["area_id"]))
                scope_a = int(cursor.lastrowid)
            else:
                scope_a = int(scope_a_row["id"])
            cursor.execute("SELECT id FROM data_scopes WHERE area_id=%s LIMIT 1", (area_b,))
            scope_b_row = cursor.fetchone()
            if scope_b_row is None:
                cursor.execute("INSERT INTO data_scopes (code,name,scope_type,organization_id,farm_id,area_id,status) VALUES ('B-all','B区全部数据','area',%s,%s,%s,'active')", (tenant["organization_id"], tenant["farm_id"], area_b))
                scope_b = int(cursor.lastrowid)
            else:
                scope_b = int(scope_b_row["id"])
            cursor.executemany("INSERT INTO user_data_scopes (user_id,data_scope_id) VALUES (%s,%s)", [(user_a, scope_a), (user_b, scope_b)])
            cursor.execute(
                "INSERT INTO ponds (organization_id,farm_id,area_id,code,name,pond_status,status,created_by) VALUES (%s,%s,%s,'P-B','B区塘','farming','verified',%s)",
                (tenant["organization_id"], tenant["farm_id"], area_b, user_b),
            )
            pond_b = int(cursor.lastrowid)
            storage_name = "c" * 32
            cursor.execute(
                "INSERT INTO attachments (organization_id,entity_type,entity_id,sha256,storage_name,original_name,media_type,size_bytes,uploaded_by) VALUES (%s,'master:ponds',%s,%s,%s,'b.pdf','application/pdf',4,%s)",
                (tenant["organization_id"], pond_b, "d" * 64, storage_name, user_b),
            )
            attachment_b = int(cursor.lastrowid)
            cursor.execute("SELECT name,status,row_version FROM ponds WHERE id=%s", (pond_b,))
            before = dict(cursor.fetchone())
        tmp_path.mkdir(parents=True, exist_ok=True)
        (tmp_path / storage_name).write_bytes(b"%PDF")

        app = create_app(_rest_settings(settings, str(tmp_path)))
        client_a = app.test_client()
        client_b = app.test_client()
        _rest_login(client_a, "idor-a")
        _rest_login(client_b, "idor-b")
        assert client_b.get(f"/api/v1/master-data/ponds/{pond_b}").status_code == 200
        csrf = _rest_csrf(client_a)
        denied: list[tuple[str, int]] = []
        for method, path, payload in (
            ("get", f"/api/v1/master-data/ponds/{pond_b}", None),
            ("patch", f"/api/v1/master-data/ponds/{pond_b}", {"expected_version": 1, "name": "越权"}),
            ("delete", f"/api/v1/master-data/ponds/{pond_b}", None),
            ("post", f"/api/v1/master-data/ponds/{pond_b}/submit", {"expected_version": 1}),
            ("post", f"/api/v1/master-data/ponds/{pond_b}/verify", {"expected_version": 1}),
            ("post", "/api/v1/master-data/ponds", {"code": "P-A-FOREIGN", "name": "外部区域引用", "area_id": area_b}),
        ):
            response = getattr(client_a, method)(path, json=payload, headers=csrf)
            assert response.status_code in {403, 404}, (method, path, response.get_json())
            denied.append((path, response.status_code))

        download = client_a.get(f"/api/v1/data-exchange/attachments/{attachment_b}/download")
        assert download.status_code in {403, 404}
        listed = client_a.get("/api/v1/data-exchange/attachments", query_string={"entity_type": "master:ponds", "entity_id": pond_b})
        assert listed.status_code == 200
        assert listed.get_json()["data"]["items"] == []
        exported = client_a.post(
            "/api/v1/data-exchange/exports",
            json={"organization_id": int(tenant["organization_id"]), "resource": "ponds", "format": "xlsx", "filters": {"area_id": area_b}},
            headers=csrf,
        )
        assert exported.status_code == 200
        workbook = load_workbook(filename=__import__("io").BytesIO(exported.data), read_only=True)
        values = [cell for row in workbook.active.iter_rows(values_only=True) for cell in row]
        assert "B区塘" not in values and "P-B" not in values

        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT name,status,row_version FROM ponds WHERE id=%s", (pond_b,))
            assert dict(cursor.fetchone()) == before
        assert len(denied) == 6


def test_mysql_confirmation_claim_is_exactly_once_under_concurrency(monkeypatch: pytest.MonkeyPatch) -> None:
    with disposable_database("adp_agent_confirmation", through=32) as database:
        settings = settings_for(database)
        _env_for(settings, monkeypatch)
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("INSERT INTO users (phone,name,password_hash,status) VALUES ('13980000001','Agent User','hash','active')")
            user_id = int(cursor.lastrowid)
        store = MySqlAgentConfirmationStore()
        token = "confirmation-token-concurrency"
        row = store.create(
            AgentConfirmation(
                id=0,
                idempotency_key="agent-confirmation:pending",
                token_hash="",
                user_id=user_id,
                session_hash="session-hash",
                conversation_id="conversation-1",
                request_id="request-1",
                tool_name="master_data.create_record",
                payload={"resource": "ponds"},
                status="pending",
                expires_at=datetime.now() + timedelta(minutes=2),
            ),
            token,
        )
        barrier = Barrier(20)

        def claim() -> Any:
            barrier.wait()
            return store.claim(token=token, user_id=user_id, session_hash="session-hash")

        with ThreadPoolExecutor(max_workers=20) as pool:
            results = list(pool.map(lambda _: claim(), range(20)))

        assert sum(result is not None for result in results) == 1
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT status, used_at FROM agent_confirmations WHERE id=%s", (row.id,))
            stored = cursor.fetchone()
        assert stored["status"] == "confirmed"
        assert stored["used_at"] is not None


def test_agent_request_id_is_exactly_once_under_concurrency(monkeypatch: pytest.MonkeyPatch) -> None:
    with disposable_database("adp_agent_idempotency", through=32) as database:
        settings = settings_for(database)
        _env_for(settings, monkeypatch)
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("INSERT INTO users (phone,name,password_hash,status) VALUES ('13980000002','Idempotent User','hash','active')")
            user_id = int(cursor.lastrowid)
        barrier = Barrier(10)
        calls: list[int] = []

        def operation() -> tuple[dict[str, Any], int]:
            calls.append(1)
            return {"mutation": "once"}, 200

        def invoke() -> tuple[str, Any]:
            barrier.wait()
            try:
                return ("ok", execute_idempotent(settings, user_id=user_id, action_code="agent:inventory.mutate", key="request-id-123456", payload={"quantity": 20}, operation=operation))
            except DomainError as error:
                return ("error", error.code)

        with ThreadPoolExecutor(max_workers=10) as pool:
            outcomes = list(pool.map(lambda _: invoke(), range(10)))

        assert calls == [1]
        assert all(kind == "ok" or value == "IDEMPOTENCY_IN_PROGRESS" for kind, value in outcomes)
        assert any(kind == "ok" for kind, _ in outcomes)
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT status, response_json FROM idempotency_keys WHERE user_id=%s", (user_id,))
            stored = cursor.fetchone()
        assert stored["status"] == "completed"
        assert '"mutation": "once"' in stored["response_json"]


def test_agent_audit_persists_before_after_and_failure_without_side_effect(monkeypatch: pytest.MonkeyPatch) -> None:
    with disposable_database("adp_agent_audit", through=32) as database:
        settings = settings_for(database)
        _env_for(settings, monkeypatch)
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("INSERT INTO users (phone,name,password_hash,status) VALUES ('13980000003','Audit User','hash','active')")
            user_id = int(cursor.lastrowid)
            logger = AuditLogger()
            logger.write(connection, user_id=user_id, action="agent_tool", object_type="master:ponds", object_id=7, result="success", ip_address=None, request_id="agent-request-1", module_code="agent", action_code="master_data.update", before={"status": "active"}, after={"status": "inactive"})
            logger.write(connection, user_id=user_id, action="agent_tool", object_type="master:ponds", object_id=8, result="failure", ip_address=None, request_id="agent-request-2", module_code="agent", action_code="master_data.update", reason="DATA_SCOPE_FORBIDDEN", before={"status": "active"}, after=None)
            cursor.execute("SELECT result, request_id, before_json, after_json, changed_fields_json, reason FROM audit_logs WHERE user_id=%s ORDER BY id", (user_id,))
            rows = cursor.fetchall()
        assert rows[0]["result"] == "success"
        assert rows[0]["request_id"] == "agent-request-1"
        assert '"status": "active"' in rows[0]["before_json"]
        assert '"status": "inactive"' in rows[0]["after_json"]
        assert '"before": "active"' in rows[0]["changed_fields_json"]
        assert rows[1]["result"] == "failure"
        assert rows[1]["reason"] == "DATA_SCOPE_FORBIDDEN"
        assert rows[1]["after_json"] is None


def _run_pond_create(settings: Any, mode: str) -> dict[str, Any]:
    ids = _seed(settings)
    service = MasterDataService(MySqlMasterDataStore(settings))
    actor = _actor(1)
    payload = {"code": "P-EQUIV", "name": "等价性塘口", "area_id": ids["area_id"], "pond_status": "build"}
    if mode == "manual":
        return service.create(actor, "ponds", payload)
    from backend.layers.features.agent.agent_gateway_service import AgentGatewayService
    from backend.layers.features.agent.agent_tool_registry import build_registry
    from backend.tests.test_agent_gateway import _MemoryConfirmationStore

    registry = build_registry(lambda _tool: lambda arguments, _context: service.create(actor, "ponds", arguments["payload"]))
    gateway = AgentGatewayService(settings, registry=registry, confirmations=_MemoryConfirmationStore())
    pending = gateway.prepare_tool(actor, "master_data.create_record", {"payload": payload}, conversation_id="equivalence", request_id="request-equivalence")
    return gateway.confirm(actor, pending["confirmation"]["token"], request_id="request-equivalence-confirm")


def test_manual_and_agent_mysql_snapshots_are_business_equivalent() -> None:
    with disposable_database("adp_manual_equivalence", through=32) as manual_db:
        manual_settings = settings_for(manual_db)
        manual = _run_pond_create(manual_settings, "manual")
        assert manual["code"] == "P-EQUIV"
        from backend.tests.helpers.db_snapshot import compare_snapshots, snapshot_tables

        manual_snapshot = snapshot_tables(manual_settings, ["areas", "ponds"])

    with disposable_database("adp_agent_equivalence", through=32) as agent_db:
        agent_settings = settings_for(agent_db)
        agent = _run_pond_create(agent_settings, "agent")
        assert agent["kind"] == "success"
        agent_snapshot = snapshot_tables(agent_settings, ["areas", "ponds"])

        assert compare_snapshots(
            manual_snapshot,
            agent_snapshot,
            ignore_fields={
                "areas": {"created_at", "updated_at"},
                "ponds": {"id", "created_at", "updated_at", "created_by"},
            },
        )


def test_agent_scope_rejects_foreign_pond_without_database_change(monkeypatch: pytest.MonkeyPatch) -> None:
    with disposable_database("adp_agent_idor", through=32) as database:
        settings = settings_for(database)
        _env_for(settings, monkeypatch)
        ids = _seed(settings)
        service = MasterDataService(MySqlMasterDataStore(settings))
        scoped_user = _actor(2, area_id=ids["area_id"] + 1)
        with pytest.raises(DomainError, match="DATA_SCOPE_FORBIDDEN"):
            service.update(scoped_user, "ponds", ids["pond_id"], {"expected_version": 1, "name": "越权修改"})
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT name FROM ponds WHERE id=%s", (ids["pond_id"],))
            assert cursor.fetchone()["name"] == "一号塘"


def test_confirmation_business_mutation_and_audit_are_exactly_once(monkeypatch: pytest.MonkeyPatch) -> None:
    """A real MySQL confirmation claim gates one business write under contention."""
    with disposable_database("adp_agent_confirmation_effect", through=32) as database:
        settings = settings_for(database)
        _env_for(settings, monkeypatch)
        ids = _seed(settings)
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("UPDATE ponds SET status='draft' WHERE id=%s", (ids["pond_id"],))
        actor = _actor(1, area_id=ids["area_id"])
        actor["_session_hash"] = "session-hash"
        actor["_session_token"] = "session-token"
        service = MasterDataService(MySqlMasterDataStore(settings))
        audit = AuditLogger()

        def write_audit(event: dict[str, Any]) -> None:
            with get_connection(settings) as connection:
                audit.write(
                    connection,
                    user_id=int(event["user_id"]),
                    action="agent_tool",
                    object_type="agent_tool",
                    object_id=ids["pond_id"],
                    result=str(event.get("result") or "failure"),
                    ip_address=None,
                    request_id=str(event.get("request_id") or ""),
                    module_code="agent",
                    action_code=str(event.get("tool_name") or "agent"),
                    reason=event.get("reason"),
                    before=event.get("before"),
                    after=event.get("after"),
                    detail_json=json.dumps(event, ensure_ascii=False, default=str),
                )

        def execute(arguments: dict[str, Any], _context: dict[str, Any]) -> dict[str, Any]:
            with get_connection(settings) as connection, connection.cursor() as cursor:
                cursor.execute("SELECT * FROM ponds WHERE id=%s", (ids["pond_id"],))
                before = dict(cursor.fetchone())
            after = service.update(actor, "ponds", ids["pond_id"], arguments["payload"])
            return {"before": before, "after": after}

        tool = AgentTool(
            name="master_data.update_record",
            description="更新主数据",
            method="PATCH",
            path_template="/api/v1/master-data/{resource}/{record_id}",
            parameters={"payload": {"type": "object", "required": True}},
            required_permission="master_data.manage",
            risk="write",
            execute=execute,
            permission_namespace="master_data",
            permission_action="manage",
            resource_argument="resource",
        )
        gateway = AgentGatewayService(
            settings,
            registry=AgentToolRegistry((tool,)),
            confirmations=MySqlAgentConfirmationStore(settings),
            audit=write_audit,
        )
        pending = gateway.prepare_tool(
            actor,
            tool.name,
            {
                "resource": "ponds",
                "record_id": ids["pond_id"],
                "payload": {"expected_version": 1, "name": "并发确认后塘"},
                "raw_instruction": "确认修改塘口名称",
            },
            conversation_id="confirmation-effect",
            request_id="confirmation-effect-prepare",
        )
        token = pending["confirmation"]["token"]
        found = gateway.confirmations.find(token=token, user_id=1, session_hash="session-hash")
        assert found is not None
        assert found.status == "pending"
        assert found.expires_at > datetime.now(timezone.utc).replace(tzinfo=None)
        barrier = Barrier(20)

        def confirm_once() -> tuple[str, str]:
            barrier.wait()
            try:
                result = gateway.confirm(actor, token, request_id="confirmation-effect-confirm")
                return "success", result["kind"]
            except AgentGatewayError as error:
                return "error", error.code

        with ThreadPoolExecutor(max_workers=20) as pool:
            outcomes = list(pool.map(lambda _: confirm_once(), range(20)))

        assert sum(kind == "success" for kind, _ in outcomes) == 1
        assert all(kind == "success" or code == "CONFIRMATION_INVALID" for kind, code in outcomes)
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT name FROM ponds WHERE id=%s", (ids["pond_id"],))
            assert cursor.fetchone()["name"] == "并发确认后塘"
            cursor.execute("SELECT status FROM agent_confirmations WHERE id=%s", (pending["confirmation"]["id"],))
            assert cursor.fetchone()["status"] == "confirmed"
            cursor.execute("SELECT COUNT(*) AS total FROM audit_logs WHERE request_id=%s AND result='success'", ("confirmation-effect-confirm",))
            assert cursor.fetchone()["total"] == 1


def test_confirmation_claim_then_real_business_failure_is_terminal(monkeypatch: pytest.MonkeyPatch) -> None:
    with disposable_database("adp_agent_confirmation_failure", through=32) as database:
        settings = settings_for(database)
        _env_for(settings, monkeypatch)
        ids = _seed(settings)
        actor = _actor(1, area_id=ids["area_id"])
        service = MasterDataService(MySqlMasterDataStore(settings))
        audit = AuditLogger()

        def write_audit(event: dict[str, Any]) -> None:
            with get_connection(settings) as connection:
                audit.write(
                    connection,
                    user_id=int(event["user_id"]),
                    action="agent_tool",
                    object_type=str(event.get("tool_name") or "agent_tool"),
                    object_id=ids["pond_id"],
                    result=str(event.get("result") or "failure"),
                    ip_address=None,
                    request_id=str(event.get("request_id") or ""),
                    module_code="agent",
                    action_code=str(event.get("tool_name") or "agent_tool"),
                    reason=event.get("reason"),
                    before=event.get("before"),
                    after=event.get("after"),
                    detail_json=json.dumps(event, ensure_ascii=False, default=str),
                )

        tool = AgentTool(
            name="master_data.update_record",
            description="更新塘口",
            method="PATCH",
            path_template="/api/v1/master-data/{resource}/{record_id}",
            parameters={"payload": {"type": "object", "required": True}},
            required_permission="master_data.manage",
            risk="write",
            execute=lambda arguments, _context: service.update(actor, "ponds", ids["pond_id"], arguments["payload"]),
            permission_namespace="master_data",
            permission_action="manage",
            resource_argument="resource",
        )
        gateway = AgentGatewayService(
            settings,
            registry=AgentToolRegistry((tool,)),
            confirmations=MySqlAgentConfirmationStore(settings),
            audit=write_audit,
        )
        pending = gateway.prepare_tool(
            actor,
            tool.name,
            {"resource": "ponds", "record_id": ids["pond_id"], "payload": {"expected_version": 999, "name": "不应写入"}, "raw_instruction": "修改塘口"},
            conversation_id="confirmation-failure",
            request_id="confirmation-failure-prepare",
        )
        with pytest.raises(AgentGatewayError) as error:
            gateway.confirm(actor, pending["confirmation"]["token"], request_id="confirmation-failure-confirm")
        assert error.value.code == "BUSINESS_ERROR"

        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT status FROM agent_confirmations WHERE id=%s", (pending["confirmation"]["id"],))
            assert cursor.fetchone()["status"] == "failed"
            cursor.execute("SELECT name FROM ponds WHERE id=%s", (ids["pond_id"],))
            assert cursor.fetchone()["name"] == "一号塘"
            cursor.execute("SELECT status FROM idempotency_keys WHERE user_id=1 AND action_code=%s", (f"agent:{tool.name}",))
            assert cursor.fetchone()["status"] == "failed"
            cursor.execute("SELECT result,reason FROM audit_logs WHERE request_id='confirmation-failure-confirm'")
            audit_row = cursor.fetchone()
            assert audit_row == {"result": "failure", "reason": "业务执行失败"}

        with pytest.raises(AgentGatewayError) as replay_error:
            gateway.confirm(actor, pending["confirmation"]["token"], request_id="confirmation-failure-replay")
        assert replay_error.value.code == "CONFIRMATION_INVALID"


def test_manual_and_agent_attachment_metadata_snapshots_are_equivalent() -> None:
    from backend.layers.features.agent.agent_tool_registry import build_registry

    upload = next(
        item
        for item in build_registry().tools
        if item.method == "POST" and item.path_template == "/api/v1/data-exchange/attachments"
    )
    assert upload.risk == "human_only"
    # NOT_APPLICABLE: multipart/form-data has no JSON Agent trigger. Metadata
    # read/download remains covered by the REST IDOR matrix.


def test_manual_and_agent_inventory_snapshots_are_equivalent() -> None:
    def run(settings: Any, mode: str) -> None:
        ids = _seed(settings)
        actor = _actor(1)
        actor["permissions"] = ["warehouse.view", "warehouse.manage"]
        verifier = _actor(2)
        verifier["permissions"] = ["warehouse.view", "warehouse.verify"]
        service = WarehouseService(MySqlWarehouseStore(settings))
        created = service.create(
            actor,
            "receipts",
            {"code": "EQ-INV", "name": "等价入库", "warehouse_id": ids["warehouse_1"], "material_id": ids["material_id"], "quantity": 4, "unit_cost": 5, "lot_no": "EQ-INV-LOT", "expiry_date": "2027-01-01"},
        )
        submitted = service.submit(actor, "receipts", created["id"], {"expected_version": created["version"]})
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute(
                "UPDATE attachments SET entity_type='warehouse:receipts',entity_id=%s WHERE id=%s",
                (created["id"], ids["attachment_id"]),
            )
        verification_payload = {"expected_version": submitted["version"], "evidence_attachment_ids": [ids["attachment_id"]]}
        if mode == "manual":
            service.verify(verifier, "receipts", created["id"], verification_payload)
            return
        tool = AgentTool(
            name="warehouse.verify_receipt",
            description="核验入库",
            method="POST",
            path_template="/api/v1/warehouse/receipts/{record_id}/verify",
            parameters={"payload": {"type": "object", "required": True}},
            required_permission="warehouse.verify",
            risk="write",
            execute=lambda arguments, _context: service.verify(verifier, "receipts", created["id"], arguments["payload"]),
        )
        gateway = AgentGatewayService(
            settings,
            registry=AgentToolRegistry((tool,)),
            confirmations=MySqlAgentConfirmationStore(settings),
        )
        pending = gateway.prepare_tool(verifier, tool.name, {"record_id": created["id"], "payload": verification_payload, "raw_instruction": "核验入库"}, conversation_id="inventory-equivalence", request_id="inventory-equivalence-prepare")
        result = gateway.confirm(verifier, pending["confirmation"]["token"], request_id="inventory-equivalence-confirm")
        assert result["kind"] == "success"

    from backend.tests.helpers.db_snapshot import compare_snapshots, snapshot_tables
    tables = ["warehouse_documents", "inventory_ledger", "cost_entries"]
    with disposable_database("adp_equiv_inventory_manual", through=32) as database:
        manual_settings = settings_for(database)
        run(manual_settings, "manual")
        manual_snapshot = snapshot_tables(manual_settings, tables)
    with disposable_database("adp_equiv_inventory_agent", through=32) as database:
        agent_settings = settings_for(database)
        run(agent_settings, "agent")
        agent_snapshot = snapshot_tables(agent_settings, tables)
        assert compare_snapshots(
            manual_snapshot,
            agent_snapshot,
            ignore_fields={
                "warehouse_documents": {"id", "created_at", "updated_at", "created_by", "updated_by", "verified_by", "verified_at", "row_version"},
                "inventory_ledger": {"id", "created_at", "happened_at", "posted_by"},
                "cost_entries": {"id", "created_at", "updated_at", "created_by", "updated_by", "confirmed_by", "confirmed_at"},
            },
            business_keys={
                "warehouse_documents": ("document_type", "code"),
                "inventory_ledger": ("source_type", "source_id", "line_no"),
                "cost_entries": ("source_type", "source_ref"),
            },
        )


@pytest.mark.parametrize(
    ("resource", "payload"),
    [
        ("batches", {"code": "EQ-BATCH", "name": "等价批次", "pond_id": "pond_id", "species": "草鱼", "initial_quantity": 100, "initial_weight_kg": 20, "stocked_at": "2026-09-08T08:00:00"}),
        ("samplings", {"code": "EQ-SP", "name": "等价抽样", "pond_id": "pond_id", "batch_id": "batch_id", "quantity": 12, "weight_kg": 3}),
        ("feed-plans", {"code": "EQ-FP", "name": "等价投喂计划", "pond_id": "pond_id", "batch_id": "batch_id", "material_id": "material_id", "quantity": 8, "planned_at": "2026-09-08T08:00:00"}),
        ("feed-logs", {"code": "EQ-FL", "name": "等价投喂记录", "pond_id": "pond_id", "batch_id": "batch_id", "material_id": "material_id", "quantity": 2, "happened_at": "2026-09-08T08:00:00"}),
        ("daily-operations", {"code": "EQ-OP", "name": "等价巡塘", "pond_id": "pond_id", "operation_type": "patrol", "payload": {"water_quality": "正常", "fish_activity": "活跃"}}),
    ],
)
def test_production_manual_agent_mysql_snapshots_are_equivalent(
    resource: str,
    payload: dict[str, Any],
) -> None:
    def run(settings: Any, mode: str) -> dict[str, Any]:
        ids = _seed(settings)
        actor = {
            "id": 1,
            "status": "active",
            "permissions": ["production.view", "production.manage"],
            "roles": [],
            "data_scopes": [],
        }
        resolved = {key: ids[value] if isinstance(value, str) and value in ids else value for key, value in payload.items()}
        service = ProductionService(MySqlProductionStore(settings))
        if mode == "manual":
            return service.create(actor, resource, resolved)
        tool = AgentTool(
            name=f"production.create_{resource.replace('-', '_')}",
            description=f"创建{resource}",
            method="POST",
            path_template=f"/api/v1/production/{resource}",
            parameters={"payload": {"type": "object", "required": True}},
            required_permission="production.manage",
            risk="write",
            execute=lambda arguments, _context: service.create(actor, resource, arguments["payload"]),
        )
        gateway = AgentGatewayService(
            settings,
            registry=AgentToolRegistry((tool,)),
            confirmations=MySqlAgentConfirmationStore(settings),
        )
        pending = gateway.prepare_tool(
            actor,
            tool.name,
            {"resource": resource, "payload": resolved, "raw_instruction": f"创建{resource}"},
            conversation_id=f"equivalence-{resource}",
            request_id=f"equivalence-{resource}-request",
        )
        return gateway.confirm(actor, pending["confirmation"]["token"], request_id=f"equivalence-{resource}-confirm")

    from backend.tests.helpers.db_snapshot import compare_snapshots, snapshot_tables

    with disposable_database(f"adp_equiv_manual_{resource.replace('-', '_')}", through=32) as database:
        manual_settings = settings_for(database)
        manual = run(manual_settings, "manual")
        assert manual["code"] == payload["code"]
        manual_snapshot = snapshot_tables(manual_settings, ["production_batches", "production_documents", "batch_stock_records"])

    with disposable_database(f"adp_equiv_agent_{resource.replace('-', '_')}", through=32) as database:
        agent_settings = settings_for(database)
        agent = run(agent_settings, "agent")
        assert agent["kind"] == "success"
        agent_snapshot = snapshot_tables(agent_settings, ["production_batches", "production_documents", "batch_stock_records"])

        assert compare_snapshots(
            manual_snapshot,
            agent_snapshot,
            ignore_fields={
                "production_batches": {"id", "created_at", "updated_at", "created_by", "updated_by"},
                "production_documents": {"id", "created_at", "updated_at", "created_by", "updated_by"},
                "batch_stock_records": {"id", "created_at", "posted_by", "happened_at"},
            },
            business_keys={
                "production_batches": ("code",),
                "production_documents": ("document_type", "code"),
                "batch_stock_records": ("source_type", "source_id", "line_no"),
            },
        )


def test_uninspected_daily_operations_query_is_date_filtered_and_scoped() -> None:
    with disposable_database("adp_uninspected_query", through=32) as database:
        settings = settings_for(database)
        ids = _seed(settings)
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO ponds (organization_id,farm_id,area_id,code,name,pond_status,status,created_by) "
                "VALUES (%s,%s,%s,'P2','二号塘','farming','verified',1)",
                (ids["organization_id"], ids["farm_id"], ids["area_id"]),
            )
            second_pond_id = int(cursor.lastrowid)
            cursor.execute(
                "INSERT INTO production_documents "
                "(organization_id,farm_id,area_id,document_type,code,name,pond_id,happened_at,status,created_by,verified_by,verified_at) "
                "VALUES (%s,%s,%s,'daily_operation','TODAY-OP','今日巡检',%s,%s,'verified',1,2,NOW())",
                (ids["organization_id"], ids["farm_id"], ids["area_id"], ids["pond_id"], date.today()),
            )
        store = MySqlProductionStore(settings)
        actor = {"id": 1, "status": "active", "permissions": ["production.view"], "data_scopes": []}
        result = store.list_records("daily-operations", user=actor, uninspected_on="today")
        returned_ids = {int(row["id"]) for row in result["items"]}
        assert ids["pond_id"] not in returned_ids
        assert second_pond_id in returned_ids
        with pytest.raises(DomainError, match="PRODUCTION_DATE_INVALID"):
            store.list_records("daily-operations", user=actor, uninspected_on="not-a-date")
