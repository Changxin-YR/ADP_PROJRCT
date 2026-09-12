from __future__ import annotations

from io import BytesIO
from typing import Any, Callable

from openpyxl import load_workbook

from backend.layers.common.db.connection import get_connection
from backend.layers.features.agent.agent_confirmation_store import MySqlAgentConfirmationStore
from backend.layers.features.agent.agent_gateway_service import AgentGatewayService
from backend.layers.features.agent.agent_tool_registry import AgentTool, AgentToolRegistry
from backend.layers.common.db.repositories.cost_store import MySqlCostStore
from backend.layers.features.cost.cost_enterprise_service import CostEnterpriseService
from backend.layers.features.data_exchange.data_exchange_service import DataExchangeService
from backend.layers.features.data_exchange.data_exchange_store import MySqlDataExchangeStore
from backend.layers.features.master_data.master_data_service import MasterDataService
from backend.layers.features.master_data.master_data_store import MySqlMasterDataStore
from backend.layers.features.production.production_service import ProductionService
from backend.layers.features.production.production_store import MySqlProductionStore
from backend.layers.features.purchase.purchase_service import PurchaseService
from backend.layers.features.purchase.purchase_store import MySqlPurchaseStore
from backend.layers.features.sales.sales_service import SalesService
from backend.layers.features.sales.sales_store import MySqlSalesStore
from backend.layers.features.warehouse.warehouse_service import WarehouseService
from backend.layers.features.warehouse.warehouse_store import MySqlWarehouseStore
from backend.tests.helpers.db_snapshot import compare_snapshots, snapshot_tables
from backend.tests.mysql_test_database import disposable_database, settings_for
from test_cost_enterprise_mysql import expense as cost_expense
from test_cost_enterprise_mysql import seed as cost_seed
from test_purchase_mysql_integration import _seed as purchase_seed
from test_sales_mysql_integration import seed as sales_seed
from test_warehouse_mysql_integration import _seed as production_seed


def _actor(user_id: int, permissions: list[str], area_id: int | None = None) -> dict[str, Any]:
    return {
        "id": user_id,
        "status": "active",
        "permissions": permissions,
        "roles": [],
        "data_scopes": [] if area_id is None else [{"scope_type": "area", "area_id": area_id}],
        "_scope_enforced": area_id is not None,
        "_session_hash": f"equivalence-session-{user_id}",
    }


def _agent_write(
    settings: Any,
    actor: dict[str, Any],
    name: str,
    permission: str,
    operation: Callable[[dict[str, Any]], dict[str, Any]],
    payload: dict[str, Any],
) -> dict[str, Any]:
    tool = AgentTool(
        name=name,
        description=name,
        method="POST",
        path_template=f"/api/v1/equivalence/{name}",
        parameters={"payload": {"type": "object", "required": True}},
        required_permission=permission,
        risk="write",
        execute=lambda arguments, _context: operation(arguments["payload"]),
    )
    gateway = AgentGatewayService(
        settings,
        registry=AgentToolRegistry((tool,)),
        confirmations=MySqlAgentConfirmationStore(settings),
    )
    pending = gateway.prepare_tool(
        actor,
        name,
        {"payload": payload, "raw_instruction": name},
        conversation_id=f"equivalence-{name}",
        request_id=f"equivalence-{name}-prepare",
    )
    result = pending if pending.get("kind") == "executed" else gateway.confirm(actor, pending["confirmation"]["token"], request_id=f"equivalence-{name}-confirm")
    return result.get("data", result)


def _business_compare(
    database_prefix: str,
    setup: Callable[[Any], tuple[dict[str, int], dict[str, Any]]],
    operation: Callable[[Any, dict[str, Any], dict[str, int]], dict[str, Any]],
    tables: list[str],
    ignore_fields: dict[str, set[str]],
) -> None:
    with disposable_database(f"{database_prefix}_manual", through=32) as database:
        settings = settings_for(database)
        ids, actor = setup(settings)
        operation(settings, {**actor, "_mode": "manual"}, ids)
        manual = snapshot_tables(settings, tables)
    with disposable_database(f"{database_prefix}_agent", through=32) as database:
        settings = settings_for(database)
        ids, actor = setup(settings)
        operation(settings, {**actor, "_mode": "agent"}, ids)
        agent = snapshot_tables(settings, tables)
    assert compare_snapshots(manual, agent, ignore_fields=ignore_fields)


def test_final_manual_agent_cost_equivalence() -> None:
    def setup(settings: Any) -> tuple[dict[str, int], dict[str, Any]]:
        ids = cost_seed(settings)
        return ids, _actor(1, ["cost.entry.manage"], ids["area_id"])

    def operation(settings: Any, actor: dict[str, Any], ids: dict[str, int]) -> dict[str, Any]:
        service = CostEnterpriseService(MySqlCostStore(settings))
        payload = cost_expense(ids, "EQ-COST")
        if actor["_mode"] == "manual":
            return service.create_expense(actor, payload)
        return _agent_write(settings, actor, "cost.create_expense", "cost.entry.manage", lambda value: service.create_expense(actor, value), payload)

    _business_compare("adp_final_equiv_cost", setup, operation, ["cost_entries"], {"cost_entries": {"id", "created_at", "updated_at"}})


def test_final_manual_agent_purchase_equivalence() -> None:
    def setup(settings: Any) -> tuple[dict[str, int], dict[str, Any]]:
        ids = purchase_seed(settings)
        return ids, _actor(1, ["purchase.view", "purchase.manage"])

    def operation(settings: Any, actor: dict[str, Any], ids: dict[str, int]) -> dict[str, Any]:
        service = PurchaseService(MySqlPurchaseStore(settings))
        payload = {"code": "EQ-PO", "name": "等价采购", "supplier_id": ids["supplier_id"], "material_id": ids["material_id"], "warehouse_id": ids["warehouse_id"], "quantity": 3, "unit_price": 5, "due_date": "2026-12-31"}
        if actor["_mode"] == "manual":
            return service.create_order(actor, payload)
        return _agent_write(settings, actor, "purchase.create_order", "purchase.manage", lambda value: service.create_order(actor, value), payload)

    _business_compare("adp_final_equiv_purchase", setup, operation, ["purchase_orders"], {"purchase_orders": {"id", "created_at", "updated_at"}})


def test_final_manual_agent_sales_equivalence() -> None:
    def setup(settings: Any) -> tuple[dict[str, int], dict[str, Any]]:
        ids = sales_seed(settings)
        return ids, _actor(1, ["sales.view", "sales.manage"])

    def operation(settings: Any, actor: dict[str, Any], ids: dict[str, int]) -> dict[str, Any]:
        service = SalesService(MySqlSalesStore(settings))
        payload = {"code": "EQ-SO", "name": "等价销售", "customer_id": ids["customer_id"], "pond_id": ids["pond_id"], "batch_id": ids["batch_id"], "species": "鲈鱼", "quantity": 2, "unit": "kg", "unit_price": 10, "sold_at": "2026-09-08", "due_date": "2026-12-31"}
        if actor["_mode"] == "manual":
            return service.create_order(actor, payload)
        return _agent_write(settings, actor, "sales.create_order", "sales.manage", lambda value: service.create_order(actor, value), payload)

    _business_compare("adp_final_equiv_sales", setup, operation, ["sales_orders"], {"sales_orders": {"id", "created_at", "updated_at"}})


def _purchase_payment_setup(settings: Any) -> tuple[dict[str, int], dict[str, Any], dict[str, Any], dict[str, Any]]:
    ids = purchase_seed(settings)
    maker = _actor(1, ["purchase.view", "purchase.manage", "finance.payable.view", "finance.payment.manage"])
    checker = _actor(2, ["purchase.view", "finance.payable.view", "finance.payment.verify"])
    service = PurchaseService(MySqlPurchaseStore(settings))
    order = service.create_order(maker, {"code": "EQ-PAY-PO", "name": "付款等价采购", "supplier_id": ids["supplier_id"], "material_id": ids["material_id"], "warehouse_id": ids["warehouse_id"], "quantity": 2, "unit_price": 5, "due_date": "2026-12-31"})
    with get_connection(settings) as connection, connection.cursor() as cursor:
        cursor.execute("INSERT INTO warehouse_documents (organization_id,farm_id,area_id,document_type,code,name,warehouse_id,material_id,purchase_order_id,quantity,unit_cost,lot_no,status,created_by) VALUES (%s,%s,%s,'receipt','EQ-PAY-REC','付款来源',%s,%s,%s,2,5,'EQ-PAY-LOT','verified',1)", (ids["organization_id"], ids["farm_id"], ids["area_id"], ids["warehouse_id"], ids["material_id"], order["id"]))
        receipt_id = int(cursor.lastrowid)
        cursor.execute("INSERT INTO purchase_payables (organization_id,purchase_order_id,source_receipt_id,supplier_id,idempotency_key,amount,due_date) VALUES (%s,%s,%s,%s,'EQ-PAYABLE',10,'2026-12-31')", (ids["organization_id"], order["id"], receipt_id, ids["supplier_id"]))
        payable_id = int(cursor.lastrowid)
    payment = service.create_payment(maker, {"code": "EQ-PAY", "name": "等价付款", "payable_id": payable_id, "amount": 10, "paid_at": "2026-09-08", "payment_method": "bank_transfer"})
    payment = service.submit_payment(maker, payment["id"], {"expected_version": payment["version"]})
    with get_connection(settings) as connection, connection.cursor() as cursor:
        cursor.execute("INSERT INTO attachments (organization_id,entity_type,entity_id,sha256,storage_name,original_name,media_type,size_bytes,uploaded_by) VALUES (%s,'purchase:payment',%s,%s,%s,'eq-pay.pdf','application/pdf',12,1)", (ids["organization_id"], payment["id"], "p" * 64, "q" * 32))
        evidence_id = int(cursor.lastrowid)
    return ids, maker, checker, {"service": service, "payment": payment, "payable_id": payable_id, "evidence_id": evidence_id}


def test_final_manual_agent_payment_equivalence() -> None:
    def run(mode: str) -> dict[str, list[dict[str, Any]]]:
        with disposable_database(f"adp_final_equiv_payment_{mode}", through=32) as database:
            settings = settings_for(database)
            _ids, _maker, checker, state = _purchase_payment_setup(settings)
            payload = {"expected_version": state["payment"]["version"], "evidence_attachment_ids": [state["evidence_id"]]}
            if mode == "manual":
                state["service"].verify_payment(checker, state["payment"]["id"], payload)
            else:
                _agent_write(settings, checker, "purchase.verify_payment", "finance.payment.verify", lambda value: state["service"].verify_payment(checker, state["payment"]["id"], value), payload)
            return snapshot_tables(settings, ["purchase_payments", "purchase_payables"])

    manual = run("manual")
    agent = run("agent")
    assert compare_snapshots(manual, agent, ignore_fields={"purchase_payments": {"id", "created_at", "updated_at", "verified_at"}, "purchase_payables": {"id", "created_at", "updated_at"}})


def _sales_receipt_setup(settings: Any) -> tuple[dict[str, int], dict[str, Any], dict[str, Any], dict[str, Any]]:
    ids = sales_seed(settings)
    maker = _actor(3, ["finance.receivable.view", "finance.receipt.manage"])
    checker = _actor(4, ["finance.receivable.view", "finance.receipt.verify"])
    with get_connection(settings) as connection, connection.cursor() as cursor:
        cursor.execute("INSERT INTO sales_orders (organization_id,farm_id,area_id,pond_id,batch_id,customer_id,code,name,species,quantity,unit,unit_price,total_amount,sold_at,due_date,status,created_by,approved_by,approved_at) VALUES (%s,%s,%s,%s,%s,%s,'EQ-REC-SO','收款等价销售','鲈鱼',2,'kg',10,20,'2026-09-08','2026-12-31','approved',1,2,NOW())", (ids["organization_id"], ids["farm_id"], ids["area_id"], ids["pond_id"], ids["batch_id"], ids["customer_id"]))
        order_id = int(cursor.lastrowid)
        cursor.execute("INSERT INTO sales_deliveries (organization_id,sales_order_id,harvest_document_id,harvest_root_id,code,name,quantity,delivered_at,status,created_by,verified_by,verified_at) VALUES (%s,%s,%s,%s,'EQ-REC-DEL','收款来源',2,'2026-09-08','verified',1,2,NOW())", (ids["organization_id"], order_id, ids["harvest_id"], ids["harvest_id"]))
        delivery_id = int(cursor.lastrowid)
        cursor.execute("INSERT INTO sales_receivables (organization_id,sales_order_id,source_delivery_id,customer_id,idempotency_key,amount,due_date) VALUES (%s,%s,%s,%s,'EQ-RECEIVABLE',20,'2026-12-31')", (ids["organization_id"], order_id, delivery_id, ids["customer_id"]))
        receivable_id = int(cursor.lastrowid)
    service = SalesService(MySqlSalesStore(settings))
    receipt = service.create_receipt(maker, {"code": "EQ-REC", "name": "等价收款", "receivable_id": receivable_id, "amount": 20, "received_at": "2026-09-08", "receipt_method": "bank_transfer"})
    receipt = service.submit_receipt(maker, receipt["id"], {"expected_version": receipt["version"]})
    with get_connection(settings) as connection, connection.cursor() as cursor:
        cursor.execute("INSERT INTO attachments (organization_id,entity_type,entity_id,sha256,storage_name,original_name,media_type,size_bytes,uploaded_by) VALUES (%s,'sales:receipt',%s,%s,%s,'eq-rec.pdf','application/pdf',12,3)", (ids["organization_id"], receipt["id"], "r" * 64, "s" * 32))
        evidence_id = int(cursor.lastrowid)
    return ids, maker, checker, {"service": service, "receipt": receipt, "evidence_id": evidence_id}


def test_final_manual_agent_receipt_equivalence() -> None:
    def run(mode: str) -> dict[str, list[dict[str, Any]]]:
        with disposable_database(f"adp_final_equiv_receipt_{mode}", through=32) as database:
            settings = settings_for(database)
            _ids, _maker, checker, state = _sales_receipt_setup(settings)
            payload = {"expected_version": state["receipt"]["version"], "evidence_attachment_ids": [state["evidence_id"]]}
            if mode == "manual":
                state["service"].verify_receipt(checker, state["receipt"]["id"], payload)
            else:
                _agent_write(settings, checker, "sales.verify_receipt", "finance.receipt.verify", lambda value: state["service"].verify_receipt(checker, state["receipt"]["id"], value), payload)
            return snapshot_tables(settings, ["sales_receipts", "sales_receivables"])

    manual = run("manual")
    agent = run("agent")
    assert compare_snapshots(manual, agent, ignore_fields={"sales_receipts": {"id", "created_at", "updated_at", "verified_at"}, "sales_receivables": {"id", "created_at", "updated_at"}})


def _purchase_return_setup(settings: Any) -> tuple[dict[str, int], dict[str, Any], dict[str, Any], dict[str, Any]]:
    ids = purchase_seed(settings)
    maker = _actor(1, ["purchase.view", "purchase.manage", "purchase.return.manage"])
    checker = _actor(2, ["purchase.view", "purchase.return.verify"])
    service = PurchaseService(MySqlPurchaseStore(settings))
    order = service.create_order(maker, {"code": "EQ-RET-PO", "name": "退货来源采购", "supplier_id": ids["supplier_id"], "material_id": ids["material_id"], "warehouse_id": ids["warehouse_id"], "quantity": 5, "unit_price": 5, "due_date": "2026-12-31"})
    with get_connection(settings) as connection, connection.cursor() as cursor:
        cursor.execute("INSERT INTO inventory_lots (organization_id,material_id,supplier_id,lot_no,unit_cost,status) VALUES (%s,%s,%s,'EQ-RET-LOT',5,'available')", (ids["organization_id"], ids["material_id"], ids["supplier_id"]))
        lot_id = int(cursor.lastrowid)
        cursor.execute("INSERT INTO warehouse_documents (organization_id,farm_id,area_id,document_type,code,name,warehouse_id,material_id,inventory_lot_id,purchase_order_id,quantity,unit_cost,lot_no,status,created_by) VALUES (%s,%s,%s,'receipt','EQ-RET-REC','退货来源',%s,%s,%s,%s,5,5,'EQ-RET-LOT','verified',1)", (ids["organization_id"], ids["farm_id"], ids["area_id"], ids["warehouse_id"], ids["material_id"], lot_id, order["id"]))
        receipt_id = int(cursor.lastrowid)
        cursor.execute("INSERT INTO inventory_ledger (organization_id,warehouse_id,material_id,inventory_lot_id,source_type,source_id,line_no,quantity_delta,unit_cost,posted_by) VALUES (%s,%s,%s,%s,'receipt',%s,1,5,5,1)", (ids["organization_id"], ids["warehouse_id"], ids["material_id"], lot_id, receipt_id))
        cursor.execute("INSERT INTO purchase_payables (organization_id,purchase_order_id,source_receipt_id,supplier_id,idempotency_key,amount,due_date) VALUES (%s,%s,%s,%s,'EQ-RET-PAYABLE',25,'2026-12-31')", (ids["organization_id"], order["id"], receipt_id, ids["supplier_id"]))
    return ids, maker, checker, {"service": service, "payload": {"code": "EQ-RET", "name": "等价采购退货", "reason": "数量复核", "source_receipt_id": receipt_id, "warehouse_id": ids["warehouse_id"], "material_id": ids["material_id"], "inventory_lot_id": lot_id, "quantity": 1}}


def test_final_manual_agent_purchase_return_equivalence() -> None:
    def run(mode: str) -> dict[str, list[dict[str, Any]]]:
        with disposable_database(f"adp_final_equiv_purchase_return_{mode}", through=32) as database:
            settings = settings_for(database)
            _ids, maker, checker, state = _purchase_return_setup(settings)
            service = state["service"]
            if mode == "manual":
                row = service.create_return(maker, state["payload"])
                row = service.submit_return(maker, row["id"], {"expected_version": row["row_version"]})
                service.verify_return(checker, row["id"], {"expected_version": row["row_version"]})
            else:
                row = _agent_write(settings, maker, "purchase.create_return", "purchase.return.manage", lambda value: service.create_return(maker, value), state["payload"])
                row = _agent_write(settings, maker, "purchase.submit_return", "purchase.return.manage", lambda value: service.submit_return(maker, row["id"], value), {"expected_version": row["row_version"]})
                _agent_write(settings, checker, "purchase.verify_return", "purchase.return.verify", lambda value: service.verify_return(checker, row["id"], value), {"expected_version": row["row_version"]})
            return snapshot_tables(settings, ["purchase_returns", "purchase_payables", "purchase_payable_adjustments", "inventory_ledger"])

    manual = run("manual")
    agent = run("agent")
    assert compare_snapshots(manual, agent, ignore_fields={"purchase_returns": {"id", "created_at", "updated_at", "verified_at"}, "purchase_payables": {"id", "created_at", "updated_at"}, "purchase_payable_adjustments": {"id", "created_at"}, "inventory_ledger": {"id", "created_at", "happened_at"}})


def _sales_return_setup(settings: Any) -> tuple[dict[str, int], dict[str, Any], dict[str, Any], dict[str, Any]]:
    ids = sales_seed(settings)
    maker = _actor(1, ["sales.view", "sales.manage", "sales.return.manage"])
    checker = _actor(2, ["sales.view", "sales.return.verify"])
    with get_connection(settings) as connection, connection.cursor() as cursor:
        cursor.execute("INSERT INTO sales_orders (organization_id,farm_id,area_id,pond_id,batch_id,customer_id,code,name,species,quantity,unit,unit_price,total_amount,sold_at,due_date,status,created_by,approved_by,approved_at) VALUES (%s,%s,%s,%s,%s,%s,'EQ-SRET-SO','退货来源销售','鲈鱼',2,'kg',10,20,'2026-09-08','2026-12-31','approved',1,2,NOW())", (ids["organization_id"], ids["farm_id"], ids["area_id"], ids["pond_id"], ids["batch_id"], ids["customer_id"]))
        order_id = int(cursor.lastrowid)
        cursor.execute("INSERT INTO sales_deliveries (organization_id,sales_order_id,harvest_document_id,harvest_root_id,code,name,quantity,delivered_at,status,created_by,verified_by,verified_at) VALUES (%s,%s,%s,%s,'EQ-SRET-DEL','销售退货来源',2,'2026-09-08','verified',1,2,NOW())", (ids["organization_id"], order_id, ids["harvest_id"], ids["harvest_id"]))
        delivery_id = int(cursor.lastrowid)
        cursor.execute("INSERT INTO sales_receivables (organization_id,sales_order_id,source_delivery_id,customer_id,idempotency_key,amount,due_date) VALUES (%s,%s,%s,%s,'EQ-SRET-RECEIVABLE',20,'2026-12-31')", (ids["organization_id"], order_id, delivery_id, ids["customer_id"]))
    return ids, maker, checker, {"service": SalesService(MySqlSalesStore(settings)), "payload": {"code": "EQ-SRET", "name": "等价销售退货", "reason": "客户退回", "source_delivery_id": delivery_id, "quantity": 1, "refund_amount": 10}}


def test_final_manual_agent_sales_return_equivalence() -> None:
    def run(mode: str) -> dict[str, list[dict[str, Any]]]:
        with disposable_database(f"adp_final_equiv_sales_return_{mode}", through=32) as database:
            settings = settings_for(database)
            _ids, maker, checker, state = _sales_return_setup(settings)
            service = state["service"]
            if mode == "manual":
                row = service.create_return(maker, state["payload"])
                row = service.submit_return(maker, row["id"], {"expected_version": row["row_version"]})
                service.verify_return(checker, row["id"], {"expected_version": row["row_version"]})
            else:
                row = _agent_write(settings, maker, "sales.create_return", "sales.return.manage", lambda value: service.create_return(maker, value), state["payload"])
                row = _agent_write(settings, maker, "sales.submit_return", "sales.return.manage", lambda value: service.submit_return(maker, row["id"], value), {"expected_version": row["row_version"]})
                _agent_write(settings, checker, "sales.verify_return", "sales.return.verify", lambda value: service.verify_return(checker, row["id"], value), {"expected_version": row["row_version"]})
            return snapshot_tables(settings, ["sales_returns", "sales_receivables", "sales_receivable_adjustments"])

    manual = run("manual")
    agent = run("agent")
    assert compare_snapshots(manual, agent, ignore_fields={"sales_returns": {"id", "created_at", "updated_at", "verified_at"}, "sales_receivables": {"id", "created_at", "updated_at"}, "sales_receivable_adjustments": {"id", "created_at"}})


def test_final_mixed_production_workflow() -> None:
    with disposable_database("adp_final_mixed_production", through=32) as database:
        settings = settings_for(database)
        ids = production_seed(settings)
        manual = _actor(1, ["production.manage", "production.daily_operations.verify"], ids["area_id"])
        agent = _actor(2, ["production.manage", "production.daily_operations.verify"], ids["area_id"])
        service = ProductionService(MySqlProductionStore(settings))
        service.create(manual, "feed-plans", {"code": "MIX-PROD-FEED", "name": "混合投喂", "pond_id": ids["pond_id"], "batch_id": ids["batch_id"], "material_id": ids["material_id"], "quantity": 2, "planned_at": "2026-09-08T08:00:00"})
        operation = {"code": "MIX-PROD-INSPECT", "name": "混合巡检", "pond_id": ids["pond_id"], "batch_id": ids["batch_id"], "operation_type": "patrol", "payload": {"water_quality": "正常", "fish_activity": "活跃"}}
        daily = service.create(manual, "daily-operations", operation)
        daily = service.submit(manual, "daily-operations", daily["id"], {"expected_version": daily["version"]})
        _agent_write(settings, agent, "production.verify_daily_operation", "production.manage", lambda value: service.verify(agent, "daily-operations", daily["id"], value), {"expected_version": daily["version"]})
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) AS total FROM production_documents WHERE code IN ('MIX-PROD-FEED','MIX-PROD-INSPECT') AND status='verified'")
            assert cursor.fetchone()["total"] == 1
            cursor.execute("SELECT status FROM production_documents WHERE code='MIX-PROD-INSPECT'")
            assert cursor.fetchone()["status"] == "verified"


def test_final_mixed_purchase_workflow() -> None:
    with disposable_database("adp_final_mixed_purchase", through=32) as database:
        settings = settings_for(database)
        ids = purchase_seed(settings)
        maker = _actor(1, ["purchase.view", "purchase.manage", "finance.payable.view", "finance.payment.manage"])
        approver = _actor(2, ["purchase.view", "purchase.verify", "warehouse.view", "warehouse.verify"])
        warehouse_maker = _actor(3, ["warehouse.view", "warehouse.manage"])
        finance = _actor(4, ["purchase.view", "finance.payable.view", "finance.payment.verify"])
        purchase = PurchaseService(MySqlPurchaseStore(settings))
        warehouse = WarehouseService(MySqlWarehouseStore(settings))
        order = purchase.create_order(maker, {"code": "MIX-PO", "name": "混合采购", "supplier_id": ids["supplier_id"], "material_id": ids["material_id"], "warehouse_id": ids["warehouse_id"], "quantity": 2, "unit_price": 5, "due_date": "2026-12-31"})
        order = purchase.submit_order(maker, order["id"], {"expected_version": order["version"]})
        order = _agent_write(settings, approver, "purchase.approve_order", "purchase.verify", lambda value: purchase.approve_order(approver, order["id"], value), {"expected_version": order["version"]})
        receipt = warehouse.create(warehouse_maker, "receipts", {"code": "MIX-REC", "name": "混合收货", "warehouse_id": ids["warehouse_id"], "material_id": ids["material_id"], "purchase_order_id": order["id"], "quantity": 2, "unit_cost": 5, "lot_no": "MIX-LOT"})
        receipt = warehouse.submit(warehouse_maker, "receipts", receipt["id"], {"expected_version": receipt["version"]})
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("UPDATE attachments SET entity_type='warehouse:receipts',entity_id=%s WHERE id=%s", (receipt["id"], ids["attachment_id"]))
        warehouse.verify(approver, "receipts", receipt["id"], {"expected_version": receipt["version"], "evidence_attachment_ids": [ids["attachment_id"]]})
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT id FROM purchase_payables WHERE source_receipt_id=%s", (receipt["id"],))
            payable_id = int(cursor.fetchone()["id"])
        payment = _agent_write(settings, maker, "purchase.create_payment", "finance.payment.manage", lambda value: purchase.create_payment(maker, value), {"code": "MIX-PAY", "name": "混合付款", "payable_id": payable_id, "amount": 10, "paid_at": "2026-09-08", "payment_method": "bank_transfer"})
        payment = purchase.submit_payment(maker, payment["id"], {"expected_version": payment["version"]})
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("INSERT INTO attachments (organization_id,entity_type,entity_id,sha256,storage_name,original_name,media_type,size_bytes,uploaded_by) VALUES (%s,'purchase:payment',%s,%s,%s,'mix-pay.pdf','application/pdf',12,1)", (ids["organization_id"], payment["id"], "m" * 64, "n" * 32))
            evidence_id = int(cursor.lastrowid)
        _agent_write(settings, finance, "purchase.verify_payment", "finance.payment.verify", lambda value: purchase.verify_payment(finance, payment["id"], value), {"expected_version": payment["version"], "evidence_attachment_ids": [evidence_id]})
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT status FROM purchase_orders WHERE id=%s", (order["id"],)); assert cursor.fetchone()["status"] == "fully_received"
            cursor.execute("SELECT status FROM purchase_payments WHERE id=%s", (payment["id"],)); assert cursor.fetchone()["status"] == "verified"
            cursor.execute("SELECT status FROM purchase_payables WHERE id=%s", (payable_id,)); assert cursor.fetchone()["status"] == "settled"


def test_final_mixed_sales_workflow() -> None:
    with disposable_database("adp_final_mixed_sales", through=32) as database:
        settings = settings_for(database)
        ids = sales_seed(settings)
        seller = _actor(1, ["sales.view", "sales.manage"])
        approver = _actor(2, ["sales.view", "sales.verify"])
        cashier = _actor(3, ["finance.receivable.view", "finance.receipt.manage"])
        checker = _actor(4, ["finance.receivable.view", "finance.receipt.verify"])
        sales = SalesService(MySqlSalesStore(settings))
        order = _agent_write(settings, seller, "sales.create_order", "sales.manage", lambda value: sales.create_order(seller, value), {"code": "MIX-SO", "name": "混合销售", "customer_id": ids["customer_id"], "pond_id": ids["pond_id"], "batch_id": ids["batch_id"], "species": "鲈鱼", "quantity": 40, "unit": "kg", "unit_price": 10, "sold_at": "2026-09-08", "due_date": "2026-12-31"})
        order = sales.submit_order(seller, order["id"], {"expected_version": order["version"]})
        order = sales.approve_order(approver, order["id"], {"expected_version": order["version"]})
        delivery = sales.create_delivery(seller, {"code": "MIX-SD", "name": "混合交付", "sales_order_id": order["id"], "harvest_document_id": ids["harvest_id"], "quantity": 40, "delivered_at": "2026-09-08"})
        delivery = sales.submit_delivery(seller, delivery["id"], {"expected_version": delivery["version"]})
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("UPDATE attachments SET entity_type='sales:delivery',entity_id=%s WHERE id=%s", (delivery["id"], ids["attachment_1"]))
        delivery = sales.verify_delivery(approver, delivery["id"], {"expected_version": delivery["version"], "evidence_attachment_ids": [ids["attachment_1"]]})
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT id FROM sales_receivables WHERE source_delivery_id=%s", (delivery["id"],)); receivable_id = int(cursor.fetchone()["id"])
        receipt = _agent_write(settings, cashier, "sales.create_receipt", "finance.receipt.manage", lambda value: sales.create_receipt(cashier, value), {"code": "MIX-SR", "name": "混合收款", "receivable_id": receivable_id, "amount": 400, "received_at": "2026-09-08", "receipt_method": "bank_transfer"})
        receipt = sales.submit_receipt(cashier, receipt["id"], {"expected_version": receipt["version"]})
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("INSERT INTO attachments (organization_id,entity_type,entity_id,sha256,storage_name,original_name,media_type,size_bytes,uploaded_by) VALUES (%s,'sales:receipt',%s,%s,%s,'mix-sr.pdf','application/pdf',12,3)", (ids["organization_id"], receipt["id"], "u" * 64, "v" * 32)); evidence_id = int(cursor.lastrowid)
        _agent_write(settings, checker, "sales.verify_receipt", "finance.receipt.verify", lambda value: sales.verify_receipt(checker, receipt["id"], value), {"expected_version": receipt["version"], "evidence_attachment_ids": [evidence_id]})
        with get_connection(settings) as connection, connection.cursor() as cursor:
            cursor.execute("SELECT status FROM sales_deliveries WHERE id=%s", (delivery["id"],)); assert cursor.fetchone()["status"] == "verified"
            cursor.execute("SELECT status FROM sales_receipts WHERE id=%s", (receipt["id"],)); assert cursor.fetchone()["status"] == "verified"
            cursor.execute("SELECT status FROM sales_receivables WHERE id=%s", (receivable_id,)); assert cursor.fetchone()["status"] == "settled"


def test_final_manual_agent_production_chain_equivalence() -> None:
    with disposable_database("adp_final_equiv_production", through=32) as database:
        settings = settings_for(database)
        ids = production_seed(settings)
        manual_actor = _actor(1, ["production.manage"], ids["area_id"])
        manual_service = ProductionService(MySqlProductionStore(settings))
        feed_payload = {"code": "EQ-MIX-FEED", "name": "混合投喂", "pond_id": ids["pond_id"], "batch_id": ids["batch_id"], "material_id": ids["material_id"], "quantity": 2, "planned_at": "2026-09-08T08:00:00"}
        manual_service.create(manual_actor, "feed-plans", feed_payload)
        manual_snapshot = snapshot_tables(settings, ["production_batches", "production_documents", "batch_stock_records"])

    with disposable_database("adp_final_equiv_production_agent", through=32) as database:
        settings = settings_for(database)
        ids = production_seed(settings)
        agent_actor = _actor(1, ["production.manage"], ids["area_id"])
        service = ProductionService(MySqlProductionStore(settings))
        feed_payload = {"code": "EQ-MIX-FEED", "name": "混合投喂", "pond_id": ids["pond_id"], "batch_id": ids["batch_id"], "material_id": ids["material_id"], "quantity": 2, "planned_at": "2026-09-08T08:00:00"}
        _agent_write(settings, agent_actor, "production.create_feed_plan", "production.manage", lambda value: service.create(agent_actor, "feed-plans", value), feed_payload)
        agent_snapshot = snapshot_tables(settings, ["production_batches", "production_documents", "batch_stock_records"])

    assert compare_snapshots(
        manual_snapshot,
        agent_snapshot,
        ignore_fields={
            "production_batches": {"id", "created_at", "updated_at", "created_by", "updated_by"},
            "production_documents": {"id", "created_at", "updated_at", "created_by", "updated_by"},
            "batch_stock_records": {"id", "created_at", "posted_by", "happened_at"},
        },
        business_keys={
            "production_batches": ("code",),
            "production_documents": ("document_type", "code"),
            "batch_stock_records": ("source_type", "source_id", "line_no"),
        },
    )


def test_final_manual_agent_export_scope_equivalence(tmp_path: Any) -> None:
    def run(settings: Any, mode: str) -> list[tuple[Any, ...]]:
        ids = production_seed(settings)
        actor = _actor(1, ["data_exchange.export", "master_data.view"], ids["area_id"])
        actor["data_scopes"][0]["organization_id"] = ids["organization_id"]
        exchange = DataExchangeService(MySqlDataExchangeStore(settings), tmp_path)
        payload = {"organization_id": ids["organization_id"], "resource": "ponds", "format": "xlsx", "filters": {"area_id": ids["area_id"]}}

        def export(value: dict[str, Any]) -> dict[str, Any]:
            content, _ = exchange.export(actor, organization_id=value["organization_id"], resource=value["resource"], file_format=value["format"], filters=value["filters"], request_id="equivalence-export")
            workbook = load_workbook(BytesIO(content), read_only=True)
            return {"rows": [tuple(row) for row in workbook.active.iter_rows(values_only=True)]}

        result = export(payload) if mode == "manual" else _agent_write(settings, actor, "data_exchange.export", "data_exchange.export", export, payload)
        return [tuple(row) for row in result["rows"]]

    with disposable_database("adp_final_equiv_export_manual", through=32) as database:
        manual = run(settings_for(database), "manual")
    with disposable_database("adp_final_equiv_export_agent", through=32) as database:
        agent = run(settings_for(database), "agent")
    assert manual == agent
